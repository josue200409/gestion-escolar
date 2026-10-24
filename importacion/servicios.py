"""Importación de grados, materias, docentes y estudiantes desde un libro de Excel."""
import re
import secrets
import unicodedata
from dataclasses import dataclass, field
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.validators import UnicodeUsernameValidator
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import IntegrityError, transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from academico.models import Asignacion, CicloLectivo, Curso, Grado, Materia, Matricula

Usuario = get_user_model()

ALFABETO_CLAVES = 'abcdefghjkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ23456789'  # sin 0/O, 1/l/I
LARGO_CLAVE = 8
_PARTICULAS = {'de', 'del', 'la', 'las', 'los', 'y', 'san', 'santa', 'van', 'von'}

# Nombres de columna aceptados (sin tildes ni mayúsculas) -> nombre interno.
_COL_GRADOS = {'grado': 'grado', 'nombre': 'grado', 'orden': 'orden'}
_COL_MATERIAS = {'materia': 'materia', 'nombre': 'materia', 'codigo': 'codigo'}
_COL_PERSONAS = {
    'nombres': 'nombres', 'nombre': 'nombres',
    'apellidos': 'apellidos', 'apellido': 'apellidos',
    'nombre completo': 'completo', 'alumno': 'completo', 'estudiante': 'completo',
    'docente': 'completo', 'maestro': 'completo', 'profesor': 'completo',
    'grado': 'grado', 'seccion': 'seccion',
    'usuario': 'usuario', 'username': 'usuario',
    'correo': 'correo', 'correo electronico': 'correo', 'email': 'correo',
    'telefono': 'telefono', 'celular': 'telefono',
}
_COL_ASIGNACIONES = {
    'docente': 'docente', 'maestro': 'docente', 'profesor': 'docente',
    'materia': 'materia', 'grado': 'grado', 'seccion': 'seccion',
}

_HOJAS = {
    'grados': {'grados', 'grado'},
    'materias': {'materias', 'materia'},
    'docentes': {'docentes', 'profesores', 'maestros'},
    'estudiantes': {'estudiantes', 'alumnos'},
    'asignaciones': {'asignaciones', 'asignacion'},
}


class _ErrorFila(Exception):
    """Problema con una fila concreta del archivo."""


@dataclass
class Resultado:
    simulacro: bool = False
    ciclo: str = ''
    errores: list = field(default_factory=list)      # (hoja, fila, mensaje)
    creados: list = field(default_factory=list)      # personas nuevas, con su contraseña temporal
    existentes: list = field(default_factory=list)   # personas que ya estaban
    conteo: dict = field(default_factory=lambda: {
        'grados': 0, 'materias': 0, 'cursos': 0, 'docentes': 0,
        'estudiantes': 0, 'matriculas': 0, 'asignaciones': 0,
    })

    @property
    def ok(self):
        return not self.errores


# ---------- utilidades de texto ----------

def _sin_acentos(texto):
    return ''.join(c for c in unicodedata.normalize('NFD', texto) if unicodedata.category(c) != 'Mn')


def _clave(texto):
    return re.sub(r'\s+', ' ', _sin_acentos(str(texto)).strip().lower())


def _texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r'\s+', ' ', str(valor)).strip()


def _nombre_propio(texto):
    return texto.title() if texto.isupper() else texto


def _ascii(texto):
    return re.sub(r'[^a-z0-9]', '', _sin_acentos(texto).lower())


def clave_aleatoria():
    return ''.join(secrets.choice(ALFABETO_CLAVES) for _ in range(LARGO_CLAVE))


# ---------- lectura del libro ----------

def _buscar_hoja(wb, tipo):
    por_clave = {_clave(n): n for n in wb.sheetnames}
    for candidato in _HOJAS[tipo]:
        if candidato in por_clave:
            return wb[por_clave[candidato]]
    return None


def _filas(hoja, alias):
    """Devuelve (columnas encontradas, [(número de fila, {columna: texto})])."""
    filas = hoja.iter_rows(values_only=True)
    encabezado = next(filas, None)
    if encabezado is None:
        return set(), []
    columnas = {}
    for i, titulo in enumerate(encabezado):
        canon = alias.get(_clave(titulo)) if titulo is not None else None
        if canon and canon not in columnas.values():
            columnas[i] = canon
    datos = []
    for n, fila in enumerate(filas, start=2):
        valores = {canon: _texto(fila[i]) if i < len(fila) else '' for i, canon in columnas.items()}
        if any(valores.values()):
            datos.append((n, valores))
    return set(columnas.values()), datos


# ---------- importador ----------

class _Importador:
    def __init__(self, wb, resultado, ciclo):
        self.wb = wb
        self.res = resultado
        self.simulacro = resultado.simulacro
        self.ciclo = ciclo
        self.grados = {_clave(g.nombre): g for g in Grado.objects.all()}
        self.materias = {_clave(m.nombre): m for m in Materia.objects.all()}
        self.ocupados = set(Usuario.objects.values_list('username', flat=True))
        self.vistos = set()
        self.cursos = {}
        self.nuevos = {}          # usuario -> datos de credenciales
        self._docentes = None

    # -- infraestructura --
    def error(self, hoja, fila, mensaje):
        self.res.errores.append((hoja, fila, mensaje))

    def _seguro(self, hoja, fila, funcion, *args):
        """Ejecuta una fila en su propio punto de guardado: si falla, solo se deshace esa fila."""
        try:
            with transaction.atomic():
                funcion(*args)
        except _ErrorFila as e:
            self.error(hoja, fila, str(e))
        except (IntegrityError, ValidationError) as e:
            self.error(hoja, fila, f'No se pudo guardar: {e}')

    def ejecutar(self):
        hojas = {tipo: _buscar_hoja(self.wb, tipo) for tipo in _HOJAS}
        if not any(hojas.values()):
            self.error('Archivo', 0, 'No encontré ninguna de las hojas esperadas: Grados, Materias, '
                       'Docentes, Estudiantes o Asignaciones. Descarga la plantilla para ver el formato.')
            return
        if hojas['grados']:
            self._catalogo('Grados', hojas['grados'], _COL_GRADOS, 'grado', self._crear_grado)
        if hojas['materias']:
            self._catalogo('Materias', hojas['materias'], _COL_MATERIAS, 'materia', self._crear_materia)
        if hojas['docentes']:
            self._personas('Docentes', hojas['docentes'], 'docente')
        if hojas['estudiantes']:
            self._personas('Estudiantes', hojas['estudiantes'], 'estudiante')
        if hojas['asignaciones']:
            self._asignaciones(hojas['asignaciones'])
        self.res.creados = [
            {**d, 'detalle': '; '.join(d['detalle'])} for d in self.nuevos.values()
        ]

    # -- grados y materias --
    def _catalogo(self, nombre_hoja, hoja, alias, columna, crear):
        cols, filas = _filas(hoja, alias)
        if columna not in cols:
            self.error(nombre_hoja, 1, f'Falta la columna «{columna.capitalize()}».')
            return
        for n, f in filas:
            self._seguro(nombre_hoja, n, crear, f)

    def _crear_grado(self, f):
        nombre = f['grado']
        if len(nombre) > 100:
            raise _ErrorFila('El nombre del grado es demasiado largo.')
        if _clave(nombre) in self.grados:
            return
        try:
            orden = int(f.get('orden') or 0) or len(self.grados) + 1
        except ValueError:
            raise _ErrorFila('El orden debe ser un número.')
        self.grados[_clave(nombre)] = Grado.objects.create(nombre=nombre, orden=orden)
        self.res.conteo['grados'] += 1

    def _crear_materia(self, f):
        nombre = f['materia']
        if len(nombre) > 100 or len(f.get('codigo', '')) > 20:
            raise _ErrorFila('El nombre de la materia o su código es demasiado largo.')
        if _clave(nombre) in self.materias:
            return
        self.materias[_clave(nombre)] = Materia.objects.create(nombre=nombre, codigo=f.get('codigo', ''))
        self.res.conteo['materias'] += 1

    # -- personas --
    def _personas(self, nombre_hoja, hoja, rol):
        cols, filas = _filas(hoja, _COL_PERSONAS)
        if not ({'nombres', 'completo'} & cols):
            self.error(nombre_hoja, 1, 'Falta la columna «Nombres» (o «Apellidos, Nombres» en una sola columna).')
            return
        if rol == 'estudiante' and 'grado' not in cols:
            self.error(nombre_hoja, 1, 'Falta la columna «Grado».')
            return
        for n, f in filas:
            self._seguro(nombre_hoja, n, self._persona, f, rol)

    def _nombres(self, f):
        nombres, apellidos = f.get('nombres', ''), f.get('apellidos', '')
        if not apellidos:
            completo = f.get('completo') or nombres
            if ',' not in completo:
                raise _ErrorFila('Falta el apellido. Usa las columnas Nombres y Apellidos, '
                                 'o escribe «Apellidos, Nombres» en una sola.')
            apellidos, nombres = [p.strip() for p in completo.split(',', 1)]
        if not nombres or not apellidos:
            raise _ErrorFila('Faltan los nombres o los apellidos.')
        if len(nombres) > 150 or len(apellidos) > 150:
            raise _ErrorFila('Los nombres o apellidos son demasiado largos.')
        return _nombre_propio(nombres), _nombre_propio(apellidos)

    def _generar_usuario(self, nombres, apellidos):
        primero = _ascii(nombres.split()[0])
        tokens = [t for t in apellidos.split() if _clave(t) not in _PARTICULAS] or apellidos.split()
        apellido = _ascii(tokens[0]) if tokens else ''
        base = f'{primero}.{apellido}' if primero and apellido else (primero or apellido or 'usuario')
        usuario, n = base, 1
        while usuario in self.ocupados:
            n += 1
            usuario = f'{base}{n}'
        return usuario

    def _buscar_existente(self, rol, usuario_pedido, nombres, apellidos):
        if usuario_pedido:
            u = Usuario.objects.filter(username=usuario_pedido).first()
            if u and u.rol != rol:
                raise _ErrorFila(f'El usuario «{u.username}» ya existe con otro rol ({u.get_rol_display()}).')
            return u
        candidatos = list(Usuario.objects.filter(
            rol=rol, first_name__iexact=nombres, last_name__iexact=apellidos)[:2])
        if len(candidatos) > 1:
            raise _ErrorFila('Hay varias personas con ese nombre. Agrega la columna «Usuario» '
                             'para indicar cuál es.')
        return candidatos[0] if candidatos else None

    def _persona(self, f, rol):
        nombres, apellidos = self._nombres(f)
        usuario_pedido = f.get('usuario', '').strip().lower()
        if usuario_pedido:
            try:
                UnicodeUsernameValidator()(usuario_pedido)
            except ValidationError:
                raise _ErrorFila(f'El usuario «{usuario_pedido}» tiene caracteres no permitidos '
                                 '(usa letras, números y . @ + - _ sin espacios).')
            if len(usuario_pedido) > 150:
                raise _ErrorFila('El usuario es demasiado largo.')
        correo = f.get('correo', '')
        if correo:
            try:
                validate_email(correo)
            except ValidationError:
                raise _ErrorFila(f'El correo «{correo}» no es válido.')
        telefono = f.get('telefono', '')
        if len(telefono) > 20:
            raise _ErrorFila('El teléfono es demasiado largo.')

        identidad = ('u', usuario_pedido) if usuario_pedido else (rol, _clave(nombres), _clave(apellidos))
        if identidad in self.vistos:
            raise _ErrorFila('Esta persona está repetida en el archivo. Si son personas distintas, '
                             'agrega la columna «Usuario» para diferenciarlas.')
        self.vistos.add(identidad)

        usuario = self._buscar_existente(rol, usuario_pedido, nombres, apellidos)
        nombre_completo = f'{nombres} {apellidos}'
        if usuario is None:
            nombre_usuario = usuario_pedido or self._generar_usuario(nombres, apellidos)
            self.ocupados.add(nombre_usuario)
            clave = None if self.simulacro else clave_aleatoria()
            usuario = Usuario.objects.create_user(
                username=nombre_usuario, password=clave, first_name=nombres, last_name=apellidos,
                email=correo, telefono=telefono, rol=rol)
            self.nuevos[nombre_usuario] = {
                'rol': usuario.get_rol_display(), 'nombre': nombre_completo,
                'usuario': nombre_usuario, 'clave': clave, 'detalle': [],
            }
            self.res.conteo['docentes' if rol == 'docente' else 'estudiantes'] += 1
        else:
            self.res.existentes.append({
                'rol': usuario.get_rol_display(), 'nombre': nombre_completo, 'usuario': usuario.username})
        if rol == 'estudiante':
            self._matricular(usuario, f)

    # -- cursos, matrículas y asignaciones --
    def _grado(self, texto):
        if not texto:
            raise _ErrorFila('Falta el grado.')
        grado = self.grados.get(_clave(texto))
        if grado is None:
            raise _ErrorFila(f'El grado «{texto}» no existe. Agrégalo en la hoja Grados.')
        return grado

    def _curso(self, grado, seccion):
        llave = (grado.id, seccion)
        if llave not in self.cursos:
            curso, creado = Curso.objects.get_or_create(ciclo=self.ciclo, grado=grado, seccion=seccion)
            self.cursos[llave] = curso
            if creado:
                self.res.conteo['cursos'] += 1
        return self.cursos[llave]

    def _matricular(self, estudiante, f):
        grado = self._grado(f.get('grado', ''))
        seccion = (f.get('seccion') or 'A').upper()[:10]
        curso = self._curso(grado, seccion)
        activa = (Matricula.objects
                  .filter(estudiante=estudiante, curso__ciclo=self.ciclo, estado=Matricula.Estado.ACTIVA)
                  .select_related('curso__grado').first())
        if activa is not None and activa.curso_id != curso.id:
            raise _ErrorFila(f'Ya está matriculado en {activa.curso}. Cámbialo desde el administrador.')
        if activa is None:
            if Matricula.objects.filter(estudiante=estudiante, curso=curso).exists():
                raise _ErrorFila('Tiene una matrícula retirada en este curso. Actívala desde el administrador.')
            Matricula.objects.create(curso=curso, estudiante=estudiante)
            self.res.conteo['matriculas'] += 1
        if estudiante.username in self.nuevos:
            self.nuevos[estudiante.username]['detalle'].append(f'{grado.nombre} {seccion}')

    def _mapa_docentes(self):
        if self._docentes is None:
            mapa = {}
            for d in Usuario.objects.filter(rol='docente'):
                nombres = {
                    d.username, f'{d.first_name} {d.last_name}',
                    f'{d.last_name} {d.first_name}', f'{d.last_name}, {d.first_name}',
                }
                for n in nombres:
                    mapa.setdefault(_clave(n), []).append(d)
            self._docentes = mapa
        return self._docentes

    def _docente(self, texto):
        if not texto:
            raise _ErrorFila('Falta el docente.')
        encontrados = {d.id: d for d in self._mapa_docentes().get(_clave(texto), [])}
        if not encontrados:
            raise _ErrorFila(f'No encuentro al docente «{texto}». Debe estar en la hoja Docentes '
                             '(o ya existir) y escribirse igual.')
        if len(encontrados) > 1:
            raise _ErrorFila(f'Hay varios docentes llamados «{texto}». Escribe su usuario.')
        return next(iter(encontrados.values()))

    def _asignaciones(self, hoja):
        cols, filas = _filas(hoja, _COL_ASIGNACIONES)
        faltan = [c for c in ('docente', 'materia', 'grado') if c not in cols]
        if faltan:
            self.error('Asignaciones', 1, 'Faltan las columnas: ' + ', '.join(f'«{c.capitalize()}»' for c in faltan) + '.')
            return
        self._docentes = None  # incluye a los docentes recién creados
        for n, f in filas:
            self._seguro('Asignaciones', n, self._asignacion, f)

    def _asignacion(self, f):
        docente = self._docente(f['docente'])
        materia = self.materias.get(_clave(f['materia']))
        if materia is None:
            raise _ErrorFila(f'La materia «{f["materia"]}» no existe. Agrégala en la hoja Materias.')
        grado = self._grado(f['grado'])
        seccion = (f.get('seccion') or 'A').upper()[:10]
        curso = self._curso(grado, seccion)
        actual = Asignacion.objects.filter(curso=curso, materia=materia).select_related('docente').first()
        if actual is None:
            Asignacion.objects.create(curso=curso, materia=materia, docente=docente)
            self.res.conteo['asignaciones'] += 1
        elif actual.docente_id != docente.id:
            raise _ErrorFila(f'{materia} en {curso} ya tiene a {actual.docente} como docente.')
        if docente.username in self.nuevos:
            self.nuevos[docente.username]['detalle'].append(f'{materia.nombre} ({grado.nombre} {seccion})')


def importar(archivo, simulacro=False):
    """Lee el libro y crea todo en una sola transacción. Si algo falla, no se guarda nada."""
    resultado = Resultado(simulacro=simulacro)
    try:
        wb = load_workbook(BytesIO(archivo.read()), data_only=True)
    except Exception:
        resultado.errores.append(('Archivo', 0, 'No se pudo leer. Sube un archivo de Excel en formato .xlsx.'))
        return resultado
    ciclo = CicloLectivo.objects.filter(activo=True).first()
    if ciclo is None:
        resultado.errores.append(('Archivo', 0, 'No hay un ciclo lectivo activo. Crea uno y márcalo '
                                  'como activo en el administrador antes de importar.'))
        return resultado
    resultado.ciclo = ciclo.nombre
    with transaction.atomic():
        _Importador(wb, resultado, ciclo).ejecutar()
        if resultado.errores or simulacro:
            transaction.set_rollback(True)
    if resultado.errores:
        resultado.creados = []
    return resultado


# ---------- libros de salida ----------

_ENCABEZADO = Font(bold=True, color='FFFFFF')
_RELLENO = PatternFill('solid', fgColor='0F3D36')


def _escribir(hoja, valores, encabezado=False):
    hoja.append(valores)
    for celda in hoja[hoja.max_row]:
        if isinstance(celda.value, str):
            celda.data_type = 's'  # nunca se interpreta como fórmula
        if encabezado:
            celda.font, celda.fill = _ENCABEZADO, _RELLENO
            celda.alignment = Alignment(vertical='center')


def _ancho(hoja, anchos):
    for letra, ancho in zip('ABCDEFGH', anchos):
        hoja.column_dimensions[letra].width = ancho


def _bytes(wb):
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def libro_credenciales(resultado):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Credenciales'
    _escribir(ws, ['Rol', 'Nombre', 'Usuario', 'Contraseña temporal', 'Curso o materias'], True)
    for c in resultado.creados:
        _escribir(ws, [c['rol'], c['nombre'], c['usuario'], c['clave'], c['detalle']])
    _ancho(ws, [14, 36, 28, 22, 50])

    ws2 = wb.create_sheet('Ya existían')
    _escribir(ws2, ['Rol', 'Nombre', 'Usuario'], True)
    for c in resultado.existentes:
        _escribir(ws2, [c['rol'], c['nombre'], c['usuario']])
    _ancho(ws2, [14, 36, 28])

    ws3 = wb.create_sheet('Resumen')
    _escribir(ws3, ['Concepto', 'Cantidad'], True)
    nombres = {'grados': 'Grados nuevos', 'materias': 'Materias nuevas', 'cursos': 'Cursos nuevos',
               'docentes': 'Docentes nuevos', 'estudiantes': 'Estudiantes nuevos',
               'matriculas': 'Matrículas nuevas', 'asignaciones': 'Asignaciones nuevas'}
    _escribir(ws3, ['Ciclo lectivo', resultado.ciclo])
    for clave, etiqueta in nombres.items():
        _escribir(ws3, [etiqueta, resultado.conteo[clave]])
    _ancho(ws3, [28, 14])
    return _bytes(wb)


def libro_plantilla():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Instrucciones'
    lineas = [
        ('Cómo llenar este archivo', True),
        ('Cada hoja es un grupo de datos. Llena solo lo que necesites; las hojas vacías se ignoran.', False),
        ('Escribe los datos desde la fila 2 de cada hoja (la fila 1 son los títulos). No cambies los títulos.', False),
        ('', False),
        ('Grados: un grado por fila (ej. Primero Básico). «Orden» es opcional.', False),
        ('Materias: una materia por fila (ej. Matemática). «Código» es opcional.', False),
        ('Docentes: Nombres y Apellidos. «Usuario», «Correo» y «Teléfono» son opcionales.', False),
        ('Estudiantes: Nombres, Apellidos, Grado y Sección. El grado debe estar en la hoja Grados.', False),
        ('Asignaciones (opcional): qué docente da qué materia en qué grado y sección.', False),
        ('', False),
        ('También puedes escribir el nombre en una sola columna como «Apellidos, Nombres».', False),
        ('Si dejas «Usuario» vacío, se crea solo con el formato nombre.apellido (ej. ana.lopez).', False),
        ('Si dos personas tienen el mismo nombre, escribe un «Usuario» distinto para cada una.', False),
        ('Los estudiantes se matriculan en el ciclo lectivo activo. No dejes datos de ejemplo en el archivo.', False),
        ('Al importar recibirás un Excel con los usuarios y las contraseñas temporales: guárdalo bien.', False),
    ]
    for texto, negrita in lineas:
        ws.append([texto])
        ws.cell(ws.max_row, 1).font = Font(bold=negrita, size=13 if negrita else 11)
    ws.column_dimensions['A'].width = 110

    hojas = {
        'Grados': ['Grado', 'Orden'],
        'Materias': ['Materia', 'Código'],
        'Docentes': ['Nombres', 'Apellidos', 'Usuario', 'Correo', 'Teléfono'],
        'Estudiantes': ['Nombres', 'Apellidos', 'Grado', 'Sección', 'Usuario', 'Correo', 'Teléfono'],
        'Asignaciones': ['Docente', 'Materia', 'Grado', 'Sección'],
    }
    for nombre, titulos in hojas.items():
        h = wb.create_sheet(nombre)
        _escribir(h, titulos, True)
        _ancho(h, [24] * len(titulos))
    return _bytes(wb)
